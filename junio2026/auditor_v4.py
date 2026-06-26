import pandas as pd
import numpy as np
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import numbers

# ============================================================
# CONFIGURACIÓN
# ============================================================
CARPETA_CSV = "archivos_csv"
ARCHIVO_SALIDA = "Reporte_Auditoria_v4.xlsx"


# ============================================================
# HELPER: Limpiar timezone de cualquier columna datetime
# ============================================================
def limpiar_timezone_fecha(df, columna):
    """Remueve timezone de una columna datetime si existe"""
    if columna in df.columns:
        df[columna] = pd.to_datetime(df[columna], errors='coerce')
        # Verificar si tiene timezone y removerlo
        if hasattr(df[columna], 'dt') and df[columna].dt.tz is not None:
            df[columna] = df[columna].dt.tz_localize(None)
    return df


# ============================================================
# 1. CARGAR CSVs
# ============================================================
def cargar_csvs():
    """Carga todos los CSVs generados por BIP"""
    print("=" * 70)
    print("🚀 PASO 1: Cargando CSVs...")
    print("=" * 70)

    archivos = {
        'inv': os.path.join(CARPETA_CSV, "core_invoices_rep.csv"),
        'dist': os.path.join(CARPETA_CSV, "distributions_rep.csv"),
        'pag': os.path.join(CARPETA_CSV, "payments_applied_rep.csv"),
        'ers': os.path.join(CARPETA_CSV, "payments_ers_rep.csv"),
        'tax': os.path.join(CARPETA_CSV, "taxes_rep.csv"),
        'wh': os.path.join(CARPETA_CSV, "withholding_rep.csv"),
    }

    dfs = {}
    for key, path in archivos.items():
        if os.path.exists(path):
            try:
                df = pd.read_csv(path, low_memory=False)
                # Limpiar nombres de columnas
                df.columns = df.columns.str.strip().str.upper()
                dfs[key] = df
                print(f"✅ {key}: {len(df):,} filas, {len(df.columns)} columnas")
            except Exception as e:
                print(f"❌ Error cargando {path}: {e}")
                dfs[key] = pd.DataFrame()
        else:
            print(f"⚠️  No encontrado: {path}")
            dfs[key] = pd.DataFrame()

    return dfs


# ============================================================
# 2. PIVOT DE IMPUESTOS (IVA desde ZX_LINES)
# ============================================================
def pivot_impuestos(df_tax):
    """Pivotea impuestos: una columna por tipo (IVA)"""
    print("\n" + "=" * 70)
    print("📊 PASO 2: Pivot de impuestos...")
    print("=" * 70)

    if df_tax.empty:
        print("⚠️  Sin datos de impuestos")
        return pd.DataFrame()

    # Asegurar INVOICE_ID consistente
    df_tax['INVOICE_ID'] = df_tax['INVOICE_ID'].astype(str)

    # Pivot: sumar IMPUESTO_USD por INVOICE_ID y CLASIFICACION_IMPUESTO
    df_pivot = df_tax.pivot_table(
        index='INVOICE_ID',
        columns='CLASIFICACION_IMPUESTO',
        values='IMPUESTO_USD',
        aggfunc='sum'
    ).fillna(0).reset_index()

    # Renombrar columnas para el reporte final
    rename_map = {
        'IVA': 'IVA_USD',
        'OTROS_TAX': 'OTROS_TAX_USD'
    }
    df_pivot = df_pivot.rename(columns=rename_map)

    # Si existe IMPUESTO_MXN, hacer otro pivot para MXN
    if 'IMPUESTO_MXN' in df_tax.columns:
        df_pivot_mxn = df_tax.pivot_table(
            index='INVOICE_ID',
            columns='CLASIFICACION_IMPUESTO',
            values='IMPUESTO_MXN',
            aggfunc='sum'
        ).fillna(0).reset_index()

        rename_map_mxn = {
            'IVA': 'IVA_MXN',
            'OTROS_TAX': 'OTROS_TAX_MXN'
        }
        df_pivot_mxn = df_pivot_mxn.rename(columns=rename_map_mxn)

        # Merge de ambos pivots
        df_pivot = pd.merge(df_pivot, df_pivot_mxn, on='INVOICE_ID', how='outer')

    print(f"✅ Impuestos pivotados: {df_pivot.shape}")
    print(f"   Columnas: {df_pivot.columns.tolist()}")

    return df_pivot


# ============================================================
# 3. PIVOT DE RETENCIONES (desde AP_INVOICE_DISTRIBUTIONS AWT)
# ============================================================
def pivot_retenciones(df_wh):
    """Pivotea retenciones: RET_ISR_4, RET_ISR_125, RET_IVA, etc."""
    print("\n" + "=" * 70)
    print("📊 PASO 3: Pivot de retenciones...")
    print("=" * 70)

    if df_wh.empty:
        print("⚠️  Sin datos de retenciones")
        return pd.DataFrame()

    df_wh['INVOICE_ID'] = df_wh['INVOICE_ID'].astype(str)

    # Pivot: sumar RETENCION_USD por INVOICE_ID y CLASIFICACION_RETENCION
    df_pivot_usd = df_wh.pivot_table(
        index='INVOICE_ID',
        columns='CLASIFICACION_RETENCION',
        values='RETENCION_USD',
        aggfunc='sum'
    ).fillna(0).reset_index()

    # Renombrar columnas
    rename_map_usd = {
        'RET_ISR_4': 'RET_ISR_4_USD',
        'RET_ISR_125': 'RET_ISR_125_USD',
        'RET_ISR_10': 'RET_ISR_10_USD',
        'RET_IVA': 'RET_IVA_USD',
        'RET_IVA_6': 'RET_IVA_6_USD',
        'RET_OTRO': 'RET_OTRO_USD'
    }
    df_pivot_usd = df_pivot_usd.rename(columns=rename_map_usd)

    # Pivot para MXN
    if 'RETENCION_MXN' in df_wh.columns:
        df_pivot_mxn = df_wh.pivot_table(
            index='INVOICE_ID',
            columns='CLASIFICACION_RETENCION',
            values='RETENCION_MXN',
            aggfunc='sum'
        ).fillna(0).reset_index()

        rename_map_mxn = {
            'RET_ISR_4': 'RET_ISR_4_MXN',
            'RET_ISR_125': 'RET_ISR_125_MXN',
            'RET_ISR_10': 'RET_ISR_10_MXN',
            'RET_IVA': 'RET_IVA_MXN',
            'RET_IVA_6': 'RET_IVA_6_MXN',
            'RET_OTRO': 'RET_OTRO_MXN'
        }
        df_pivot_mxn = df_pivot_mxn.rename(columns=rename_map_mxn)

        df_pivot = pd.merge(df_pivot_usd, df_pivot_mxn, on='INVOICE_ID', how='outer')
    else:
        df_pivot = df_pivot_usd

    print(f"✅ Retenciones pivotadas: {df_pivot.shape}")
    print(f"   Columnas: {df_pivot.columns.tolist()}")

    return df_pivot


# ============================================================
# 4. PIVOT DE DISTRIBUCIONES (cuentas contables)
# ============================================================
def pivot_distribuciones(df_dist):
    """Pivotea distribuciones: una fila por factura con cuenta contable principal"""
    print("\n" + "=" * 70)
    print("📊 PASO 4: Pivot de distribuciones...")
    print("=" * 70)

    if df_dist.empty:
        print("⚠️  Sin datos de distribuciones")
        return pd.DataFrame()

    df_dist['INVOICE_ID'] = df_dist['INVOICE_ID'].astype(str)

    # Para cada factura, tomar la primera distribución (la principal)
    df_dist_sorted = df_dist.sort_values(['INVOICE_ID', 'DISTRIBUTION_LINE_NUMBER'])
    df_pivot = df_dist_sorted.groupby('INVOICE_ID').first().reset_index()

    # Seleccionar solo las columnas necesarias
    columnas_necesarias = [
        'INVOICE_ID',
        'CUENTA_CONTABLE_GASTO',
        'CC_EMPRESA',
        'CC_DEPARTAMENTO',
        'CC_CUENTA',
        'DESCRIPTION'
    ]

    df_pivot = df_pivot[[col for col in columnas_necesarias if col in df_pivot.columns]]

    print(f"✅ Distribuciones pivotadas: {df_pivot.shape}")
    print(f"   Columnas: {df_pivot.columns.tolist()}")

    return df_pivot


# ============================================================
# 5. AGREGAR PAGOS (sumar por factura) - SIN TIMEZONE
# ============================================================
def agregar_pagos(df_pag):
    """Agrega pagos: suma IMPORTE_PAGO por factura"""
    print("\n" + "=" * 70)
    print("📊 PASO 5: Agregando pagos...")
    print("=" * 70)

    if df_pag.empty:
        print("⚠️  Sin datos de pagos")
        return pd.DataFrame()

    df_pag['INVOICE_ID'] = df_pag['INVOICE_ID'].astype(str)

    # 🔧 CRÍTICO: Limpiar timezone ANTES de cualquier operación
    for col_fecha in ['FECHA_APLICACION', 'FECHA_PAGO']:
        if col_fecha in df_pag.columns:
            df_pag = limpiar_timezone_fecha(df_pag, col_fecha)

    # Agregar pagos por factura
    df_agg = df_pag.groupby('INVOICE_ID').agg({
        'NUM_PAGO': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'BANCO_CUENTA': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'DESCRIPTION': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'IMPORTE_PAGO_USD': 'sum',
        'IMPORTE_PAGO_MXN': 'sum',
        'TIPO_CAMBIO_PAGO': 'mean',
        'FECHA_APLICACION': 'max',
        'FECHA_PAGO': 'max'
    }).reset_index()

    # 🔧 Asegurar que las fechas agregadas no tengan timezone
    for col_fecha in ['FECHA_APLICACION', 'FECHA_PAGO']:
        if col_fecha in df_agg.columns:
            df_agg = limpiar_timezone_fecha(df_agg, col_fecha)

    print(f"✅ Pagos agregados: {df_agg.shape}")
    print(f"   Columnas: {df_agg.columns.tolist()}")

    return df_agg


# ============================================================
# 6. AGREGAR ERS/PREPAYMENTS (sumar adelantos por factura) - SIN TIMEZONE
# ============================================================
def agregar_ers(df_ers):
    """Agrega ERS/prepagos: suma ADELANTO por factura"""
    print("\n" + "=" * 70)
    print("📊 PASO 6: Agregando ERS/prepagos...")
    print("=" * 70)

    if df_ers.empty:
        print("⚠️  Sin datos de ERS/prepagos")
        return pd.DataFrame()

    df_ers['INVOICE_ID'] = df_ers['INVOICE_ID'].astype(str)

    # 🔧 CRÍTICO: Limpiar timezone ANTES de cualquier operación
    if 'FECHA_APLICACION' in df_ers.columns:
        df_ers = limpiar_timezone_fecha(df_ers, 'FECHA_APLICACION')

    # Agregar por factura
    df_agg = df_ers.groupby('INVOICE_ID').agg({
        'TIPO_ORIGEN': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'ERS': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'ORDEN_COMPRA': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'FACTURA_ADELANTO': lambda x: ', '.join(x.dropna().astype(str).unique()),
        'ADELANTO_USD': 'sum',
        'ADELANTO_MXN': 'sum',
        'TIPO_CAMBIO': 'mean',
        'FECHA_APLICACION': 'max'
    }).reset_index()

    # 🔧 Asegurar que la fecha agregada no tenga timezone
    if 'FECHA_APLICACION' in df_agg.columns:
        df_agg = limpiar_timezone_fecha(df_agg, 'FECHA_APLICACION')

    # Fallback: si MONEDA = 'MXN' y ADELANTO_MXN está vacío, usar ADELANTO_USD
    if 'MONEDA' in df_ers.columns:
        moneda_map = df_ers.groupby('INVOICE_ID')['MONEDA'].first().reset_index()
        df_agg = pd.merge(df_agg, moneda_map, on='INVOICE_ID', how='left')

    # Aplicar fallback
    mask_mxn = (df_agg.get('MONEDA') == 'MXN') & (df_agg['ADELANTO_MXN'].isna())
    df_agg.loc[mask_mxn, 'ADELANTO_MXN'] = df_agg.loc[mask_mxn, 'ADELANTO_USD']

    print(f"✅ ERS/prepagos agregados: {df_agg.shape}")
    print(f"   Columnas: {df_agg.columns.tolist()}")

    return df_agg


# ============================================================
# 7. MERGE MAESTRO
# ============================================================
def merge_maestro(df_inv, df_impuestos, df_retenciones, df_dist, df_pag, df_ers):
    """Merge maestro de todos los dataframes por INVOICE_ID"""
    print("\n" + "=" * 70)
    print("🔗 PASO 7: Merge maestro...")
    print("=" * 70)

    df_inv['INVOICE_ID'] = df_inv['INVOICE_ID'].astype(str)

    # Merge con impuestos
    if not df_impuestos.empty:
        df_result = pd.merge(df_inv, df_impuestos, on='INVOICE_ID', how='left')
        print(f"✅ Merge con impuestos: {df_result.shape}")
    else:
        df_result = df_inv.copy()

    # Merge con retenciones
    if not df_retenciones.empty:
        df_result = pd.merge(df_result, df_retenciones, on='INVOICE_ID', how='left')
        print(f"✅ Merge con retenciones: {df_result.shape}")

    # Merge con distribuciones
    if not df_dist.empty:
        df_result = pd.merge(df_result, df_dist, on='INVOICE_ID', how='left')
        print(f"✅ Merge con distribuciones: {df_result.shape}")

    # Merge con pagos
    if not df_pag.empty:
        df_result = pd.merge(df_result, df_pag, on='INVOICE_ID', how='left')
        print(f"✅ Merge con pagos: {df_result.shape}")

    # Merge con ERS/prepagos
    if not df_ers.empty:
        df_result = pd.merge(df_result, df_ers, on='INVOICE_ID', how='left')
        print(f"✅ Merge con ERS/prepagos: {df_result.shape}")

    print(f"\n📊 Resultado final: {df_result.shape}")
    print(f"   Columnas: {len(df_result.columns)}")

    return df_result


# ============================================================
# 8. CALCULAR CAMPOS DERIVADOS
# ============================================================
def calcular_campos(df):
    """Calcula SUBTOTAL, SALDO, SEMANA, MES, etc."""
    print("\n" + "=" * 70)
    print("🧮 PASO 8: Calculando campos derivados...")
    print("=" * 70)

    # Convertir fechas
    for col_fecha in ['FECHA', 'TERMS_DATE']:
        if col_fecha in df.columns:
            df[col_fecha] = pd.to_datetime(df[col_fecha], errors='coerce')

    # Calcular SUBTOTAL = IMPORTE - IVA - RETENCIONES
    if 'IMPORTE' in df.columns:
        df['SUBTOTAL'] = df['IMPORTE'].fillna(0)

        # Restar IVA
        if 'IVA_USD' in df.columns:
            df['SUBTOTAL'] = df['SUBTOTAL'] - df['IVA_USD'].fillna(0)

        # Restar retenciones
        for col_ret in ['RET_ISR_4_USD', 'RET_ISR_125_USD', 'RET_ISR_10_USD', 'RET_IVA_USD', 'RET_OTRO_USD']:
            if col_ret in df.columns:
                df['SUBTOTAL'] = df['SUBTOTAL'] - df[col_ret].fillna(0)

        print(f"✅ SUBTOTAL calculado")

    # Calcular SALDO = IMPORTE - IMPORTE_PAGO - ADELANTO
    if 'IMPORTE' in df.columns:
        df['SALDO'] = df['IMPORTE'].fillna(0)

        if 'IMPORTE_PAGO_USD' in df.columns:
            df['SALDO'] = df['SALDO'] - df['IMPORTE_PAGO_USD'].fillna(0)

        if 'ADELANTO_USD' in df.columns:
            df['SALDO'] = df['SALDO'] - df['ADELANTO_USD'].fillna(0)

        print(f"✅ SALDO calculado")

    # Calcular SEMANA y MES
    if 'FECHA' in df.columns:
        df['SEMANA'] = df['FECHA'].dt.isocalendar().week
        df['MES'] = df['FECHA'].dt.month
        print(f"✅ SEMANA y MES calculados")

    # Calcular SEMANA_PAGO y MES_PAGO
    if 'FECHA_PAGO' in df.columns:
        df['SEMANA_PAGO'] = df['FECHA_PAGO'].dt.isocalendar().week
        df['MES_PAGO'] = df['FECHA_PAGO'].dt.month
        print(f"✅ SEMANA_PAGO y MES_PAGO calculados")

    # Calcular IMPORTE_BRUTO (igual a IMPORTE)
    if 'IMPORTE' in df.columns:
        df['IMPORTE_BRUTO'] = df['IMPORTE']
        print(f"✅ IMPORTE_BRUTO calculado")

    # Calcular total retenciones
    cols_ret = ['RET_ISR_4_USD', 'RET_ISR_125_USD', 'RET_ISR_10_USD', 'RET_IVA_USD', 'RET_OTRO_USD']
    cols_ret_existentes = [col for col in cols_ret if col in df.columns]
    if cols_ret_existentes:
        df['TOTAL_RETENCIONES'] = df[cols_ret_existentes].sum(axis=1)
        print(f"✅ TOTAL_RETENCIONES calculado")

    print(f"\n📊 Campos calculados: {df.shape}")

    return df


# ============================================================
# 9. CONVERTIR FECHAS A FORMATO dd/mm/aaaa (SIN TIMEZONE)
# ============================================================
def convertir_fechas(df):
    """Convierte todas las columnas de fecha a datetime sin timezone"""
    print("\n" + "=" * 70)
    print("📅 PASO 9: Convirtiendo fechas...")
    print("=" * 70)

    columnas_fecha = ['FECHA', 'TERMS_DATE', 'FECHA_APLICACION', 'FECHA_PAGO']

    for col in columnas_fecha:
        if col in df.columns:
            # 🔧 CRÍTICO: Usar la función helper para limpiar timezone
            df = limpiar_timezone_fecha(df, col)

            # Contar conversiones exitosas
            exitosas = df[col].notna().sum()
            print(f"✅ {col}: {exitosas:,} fechas convertidas")

    return df


# ============================================================
# 10. SELECCIONAR Y ORDENAR COLUMNAS FINALES
# ============================================================
# ============================================================
# 10. SELECCIONAR Y ORDENAR COLUMNAS FINALES
# ============================================================
def seleccionar_columnas_finales(df):
    """Selecciona y ordena las columnas según el Excel de Sandra"""
    print("\n" + "=" * 70)
    print("📋 PASO 10: Seleccionando columnas finales...")
    print("=" * 70)

    # 🔧 PASO A: Resolver duplicados ANTES de seleccionar
    # Para campos duplicados (_x, _y), priorizar la versión de core_invoices (_x)
    duplicados_a_resolver = {
        'MONEDA_x': 'MONEDA',
        'MONEDA_y': None,  # Eliminar
        'ERS_x': 'ERS',
        'ERS_y': None,  # Eliminar
        'ORDEN_COMPRA_x': 'ORDEN_COMPRA',
        'ORDEN_COMPRA_y': None,  # Eliminar
        'DESCRIPTION_x': 'DESCRIPTION',
        'DESCRIPTION_y': None,  # Eliminar
        'FECHA_APLICACION_x': 'FECHA_APLICACION',
        'FECHA_APLICACION_y': None,  # Eliminar
    }

    for col_actual, col_nueva in duplicados_a_resolver.items():
        if col_actual in df.columns:
            if col_nueva is None:
                # Eliminar la columna duplicada
                df = df.drop(columns=[col_actual])
                print(f"   🗑️  Eliminando duplicado: {col_actual}")
            else:
                # Renombrar la columna _x a su nombre original
                df = df.rename(columns={col_actual: col_nueva})
                print(f"   🔄 Renombrando: {col_actual} → {col_nueva}")
                # Eliminar la versión _y si existe
                col_y = col_actual.replace('_x', '_y')
                if col_y in df.columns:
                    df = df.drop(columns=[col_y])
                    print(f"   🗑️  Eliminando duplicado: {col_y}")

    # 🔧 PASO B: Eliminar IDs internos y campos no requeridos
    columnas_eliminar = [
        'VENDOR_ID',
        'ORG_ID',
        'TERMS_ID',
        'INVOICE_ID',
        'TERMS_DATE',
        'CC_EMPRESA',
        'CC_CUENTA',
        'TIPO_CAMBIO',  # Redundante
        'VOUCHER_NUM',  # No requerido por Sandra
    ]

    columnas_a_eliminar = [col for col in columnas_eliminar if col in df.columns]
    if columnas_a_eliminar:
        df = df.drop(columns=columnas_a_eliminar)
        print(f"   🗑️  Columnas eliminadas: {columnas_a_eliminar}")

    # 🔧 PASO C: Orden de columnas según Sandra-revision.xlsx
    columnas_orden = [
        # Datos de factura
        'EMPRESA',
        'NUM_PROV',
        'PROVEEDOR',
        'RFC',
        'FACTURA',
        'ERS',
        'TIPO',
        'UUID',
        'ORIGEN',
        'ESTATUS_APROBACION',

        # Cuentas contables
        'CUENTA_CONTABLE_GASTO',
        'CC_DEPARTAMENTO',

        # Fechas
        'FECHA',
        'FECHA_PAGO',
        'FECHA_APLICACION',

        # Períodos
        'SEMANA',
        'MES',
        'SEMANA_PAGO',
        'MES_PAGO',

        # Pagos
        'NUM_PAGO',
        'DESCRIPTION',
        'BANCO_CUENTA',

        # Montos USD
        'ADELANTO_USD',
        'IMPORTE_PAGO_USD',
        'SUBTOTAL',
        'IVA_USD',
        'RET_ISR_4_USD',
        'RET_ISR_125_USD',
        'RET_ISR_10_USD',
        'RET_IVA_USD',
        'RET_OTRO_USD',
        'TOTAL_RETENCIONES',
        'IMPORTE_BRUTO',
        'SALDO',

        # Montos MXN
        'ADELANTO_MXN',
        'IMPORTE_PAGO_MXN',
        'IVA_MXN',
        'RET_ISR_4_MXN',
        'RET_ISR_125_MXN',
        'RET_ISR_10_MXN',
        'RET_IVA_MXN',
        'RET_OTRO_MXN',

        # Tipos de cambio y moneda
        'TIPO_CAMBIO_FACTURA',
        'TIPO_CAMBIO_PAGO',
        'MONEDA',

        # Campos adicionales
        'IMPORTE',
        'TOTAL_TAX_AMOUNT',
        'ORDEN_COMPRA',
        'FACTURA_ADELANTO',
        'TIPO_ORIGEN',
        'DUE_DAYS',
        'ORGANIZATION_TYPE_LOOKUP_CODE'
    ]

    # Filtrar solo las columnas que existen en el dataframe
    columnas_existentes = [col for col in columnas_orden if col in df.columns]

    # Agregar columnas adicionales que no estaban en la lista (por si acaso)
    columnas_adicionales = [col for col in df.columns if col not in columnas_orden]
    if columnas_adicionales:
        print(f"   ⚠️  Columnas adicionales no esperadas: {columnas_adicionales}")

    columnas_finales = columnas_existentes + columnas_adicionales

    df_final = df[columnas_finales].copy()

    print(f"\n✅ Columnas seleccionadas: {len(columnas_finales)}")
    print(f"   Lista completa: {columnas_finales}")

    # Renombrar columnas para el reporte final (más amigables)
    rename_map = {
        'ESTATUS_APROBACION': 'ESTATUS',
        'ADELANTO_USD': 'ADELANTO',
        'IMPORTE_PAGO_USD': 'IMPORTE_PAGO',
        'IVA_USD': 'IVA',
        'RET_ISR_4_USD': 'RETENCION_ISR_4',
        'RET_ISR_125_USD': 'RETENCION_ISR_125',
        'RET_ISR_10_USD': 'RETENCION_ISR_10',
        'RET_IVA_USD': 'RETENCION_IVA',
        'RET_OTRO_USD': 'RETENCION_OTRO'
    }

    # Aplicar solo los renombres que existen
    rename_aplicar = {k: v for k, v in rename_map.items() if k in df_final.columns}
    df_final = df_final.rename(columns=rename_aplicar)

    print(f"✅ Columnas renombradas: {len(rename_aplicar)}")

    return df_final


# ============================================================
# 11. EXPORTAR A EXCEL CON FORMATOS
# ============================================================
def exportar_excel(df, archivo_salida):
    """Exporta a Excel con formatos de fecha y moneda"""
    print("\n" + "=" * 70)
    print("💾 PASO 11: Exportando a Excel...")
    print("=" * 70)

    # Reemplazar NaN con None para Excel
    df = df.replace({np.nan: None})

    # Exportar a Excel
    df.to_excel(archivo_salida, index=False, sheet_name='Reporte_Auditoria')

    print(f"✅ Excel generado: {archivo_salida}")
    print(f"   Filas: {len(df):,}")
    print(f"   Columnas: {len(df.columns)}")

    # Aplicar formatos con openpyxl
    print("\n📝 Aplicando formatos...")
    wb = load_workbook(archivo_salida)
    ws = wb['Reporte_Auditoria']

    # Obtener índices de columnas
    headers = {cell.value: cell.column for cell in ws[1]}

    # Formato de fecha: dd/mm/yyyy
    formato_fecha = 'dd/mm/yyyy'
    columnas_fecha = ['FECHA', 'FECHA_PAGO', 'FECHA_APLICACION']

    for col_nombre in columnas_fecha:
        if col_nombre in headers:
            col_idx = headers[col_nombre]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    cell.number_format = formato_fecha
            print(f"✅ Formato fecha aplicado a {col_nombre}")

    # Formato de moneda: #,##0.00
    formato_moneda = '#,##0.00'
    columnas_moneda = [
        'IMPORTE', 'ADELANTO', 'IMPORTE_PAGO', 'SUBTOTAL', 'IVA',
        'RETENCION_ISR_4', 'RETENCION_ISR_125', 'RETENCION_ISR_10',
        'RETENCION_IVA', 'RETENCION_OTRO', 'TOTAL_RETENCIONES',
        'IMPORTE_BRUTO', 'SALDO', 'TOTAL_TAX_AMOUNT',
        'ADELANTO_MXN', 'IMPORTE_PAGO_MXN', 'IVA_MXN',
        'RET_ISR_4_MXN', 'RET_ISR_125_MXN', 'RET_ISR_10_MXN',
        'RET_IVA_MXN', 'RET_OTRO_MXN'
    ]

    for col_nombre in columnas_moneda:
        if col_nombre in headers:
            col_idx = headers[col_nombre]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    cell.number_format = formato_moneda
            print(f"✅ Formato moneda aplicado a {col_nombre}")

    # Formato de tipo de cambio: #,##0.0000
    formato_tc = '#,##0.0000'
    columnas_tc = ['TIPO_CAMBIO_FACTURA', 'TIPO_CAMBIO_PAGO']

    for col_nombre in columnas_tc:
        if col_nombre in headers:
            col_idx = headers[col_nombre]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    cell.number_format = formato_tc
            print(f"✅ Formato tipo cambio aplicado a {col_nombre}")

    # Ajustar ancho de columnas automáticamente
    print("\n📏 Ajustando ancho de columnas...")
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)  # Máximo 50 caracteres
        ws.column_dimensions[column_letter].width = adjusted_width

    # Guardar cambios
    wb.save(archivo_salida)
    print(f"✅ Formatos aplicados y guardados")

    return archivo_salida


# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print(f"\n{'=' * 70}")
    print(f"🔍 GENERADOR DE REPORTE DE AUDITORÍA v4")
    print(f"📅 Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"{'=' * 70}\n")

    # Paso 1: Cargar CSVs
    dfs = cargar_csvs()

    # Validar que tenemos facturas
    if dfs['inv'].empty:
        print("\n❌ ERROR: No hay datos de facturas. Deteniendo proceso.")
        exit(1)

    # Paso 2: Pivot de impuestos
    df_impuestos = pivot_impuestos(dfs['tax'])

    # Paso 3: Pivot de retenciones
    df_retenciones = pivot_retenciones(dfs['wh'])

    # Paso 4: Pivot de distribuciones
    df_distribuciones = pivot_distribuciones(dfs['dist'])

    # Paso 5: Agregar pagos
    df_pagos = agregar_pagos(dfs['pag'])

    # Paso 6: Agregar ERS/prepagos
    df_ers_agg = agregar_ers(dfs['ers'])

    # Paso 7: Merge maestro
    df_maestro = merge_maestro(
        dfs['inv'],
        df_impuestos,
        df_retenciones,
        df_distribuciones,
        df_pagos,
        df_ers_agg
    )

    # Paso 8: Calcular campos derivados
    df_final = calcular_campos(df_maestro)

    # Paso 9: Convertir fechas
    df_fechas = convertir_fechas(df_final)

    # Paso 10: Seleccionar columnas finales
    df_columnas = seleccionar_columnas_finales(df_fechas)

    # Paso 11: Exportar a Excel
    archivo_final = exportar_excel(df_columnas, ARCHIVO_SALIDA)

    print("\n" + "=" * 70)
    print("🎉 PROCESO COMPLETADO EXITOSAMENTE")
    print("=" * 70)
    print(f"📁 Archivo generado: {archivo_final}")
    print(f"📊 Total filas: {len(df_columnas):,}")
    print(f"📋 Total columnas: {len(df_columnas.columns)}")
    print("=" * 70)